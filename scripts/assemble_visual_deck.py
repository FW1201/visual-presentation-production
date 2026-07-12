#!/usr/bin/env python3
"""Validate approved 16:9 slide images and assemble optional PPTX/PDF outputs."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from PIL import Image, ImageOps
except ImportError as error:
    raise SystemExit("Pillow is required: python3 -m pip install pillow") from error


def load_manifest(path: Path) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise SystemExit(f"Cannot read manifest: {error}") from error


def validate_manifest(data: dict) -> list[str]:
    errors: list[str] = []
    count, slides = data.get("slide_count"), data.get("slides")
    if not isinstance(count, int) or count < 2:
        errors.append("slide_count must be an integer of at least 2")
    if data.get("width_px") != 1920 or data.get("height_px") != 1080:
        errors.append("this version requires 1920x1080 output")
    if data.get("sample_approved") is not True:
        errors.append("sample_approved must be true")
    if data.get("all_text_frozen") is not True:
        errors.append("all_text_frozen must be true")
    if not isinstance(slides, list) or len(slides) != count:
        errors.append("slides must contain exactly slide_count rows")
        return errors
    ids = [str(slide.get("id", "")) for slide in slides]
    if len(ids) != len(set(ids)) or any(not item for item in ids):
        errors.append("slide ids must be present and unique")
    if any(slide.get("approved") is not True for slide in slides):
        errors.append("every slide must be approved")
    return errors


def normalize(source: Path, destination: Path) -> None:
    with Image.open(source) as raw:
        raw.load()
        if abs(raw.width / raw.height - 16 / 9) > 0.02:
            raise ValueError(f"{source.name} is not 16:9")
        image = raw.convert("RGB")
    image = ImageOps.contain(image, (1920, 1080), Image.Resampling.LANCZOS)
    canvas = Image.new("RGB", (1920, 1080), "white")
    canvas.paste(image, ((1920 - image.width) // 2, (1080 - image.height) // 2))
    destination.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(destination, "PNG", optimize=True)


def write_pdf(images: list[Path], destination: Path) -> None:
    pages = [Image.open(path).convert("RGB") for path in images]
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        pages[0].save(destination, "PDF", save_all=True, append_images=pages[1:], resolution=150)
    finally:
        for page in pages:
            page.close()


def write_pptx(images: list[Path], destination: Path) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as error:
        raise RuntimeError("python-pptx is required for PPTX assembly") from error
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    for image in images:
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_picture(str(image), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(destination)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    parser.add_argument("--pptx", type=Path)
    parser.add_argument("--pdf", type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    data = load_manifest(args.manifest)
    errors = validate_manifest(data)
    final_images: list[Path] = []
    if not errors:
        for slide in data["slides"]:
            identifier = str(slide["id"]).zfill(2)
            source = args.input_dir / f"slide_{identifier}.png"
            destination = args.out_dir / f"slide_{identifier}.png"
            if not source.exists():
                errors.append(f"missing slide_{identifier}.png")
                continue
            try:
                normalize(source, destination)
                final_images.append(destination)
            except (OSError, ValueError) as error:
                errors.append(str(error))
    if not errors:
        try:
            if args.pdf:
                write_pdf(final_images, args.pdf)
            if args.pptx:
                write_pptx(final_images, args.pptx)
        except RuntimeError as error:
            errors.append(str(error))
    report_path = args.report or args.out_dir / "qa-report.json"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps({"valid": not errors, "errors": errors}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if errors:
        print("Validation failed:\n" + "\n".join(f"- {item}" for item in errors))
        return 1
    print(f"Prepared {len(final_images)} slides in {args.out_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
